"""Export filtered data to CSV and Excel."""

from __future__ import annotations

import csv
from pathlib import Path

from .config import DB_PATH, DATA_DIR
from .filters import (
    PERFORMANCE_SELECT,
    SHOW_SELECT,
    SONG_SELECT,
    SearchFilters,
    build_performance_where,
    build_show_where,
    build_song_where,
)
from .schema import format_duration, get_connection

EXPORT_TYPES = ("performances", "shows", "songs")

PERFORMANCE_COLUMNS = [
    ("showdate", "Show Date"),
    ("song", "Song"),
    ("set_name", "Set"),
    ("position", "Position"),
    ("tracktime", "Length"),
    ("gap", "Gap"),
    ("venue", "Venue"),
    ("city", "City"),
    ("state", "State"),
    ("country", "Country"),
    ("tour_name", "Tour"),
    ("transition", "Transition"),
    ("footnote", "Notes/Guests"),
    ("is_jam", "Jam"),
    ("is_jamchart", "Jam Chart"),
    ("jamchart_description", "Jam Chart Notes"),
    ("soundcheck", "Soundcheck"),
    ("setlist_notes", "Setlist Notes"),
]

SHOW_COLUMNS = [
    ("showdate", "Show Date"),
    ("venue", "Venue"),
    ("city", "City"),
    ("state", "State"),
    ("country", "Country"),
    ("tour_name", "Tour"),
    ("soundcheck", "Soundcheck"),
    ("setlist_notes", "Setlist Notes"),
    ("permalink", "Permalink"),
]

SONG_COLUMNS = [
    ("song", "Song"),
    ("slug", "Slug"),
    ("times_played", "Times Played"),
    ("debut", "Debut"),
    ("last_played", "Last Played"),
    ("gap", "Gap"),
]


def _normalize_row(row: dict, export_type: str) -> dict:
    data = dict(row)
    if export_type == "performances":
        data["tracktime"] = format_duration(data.get("tracktime_seconds")) or data.get(
            "tracktime_raw", ""
        )
        data["is_jam"] = "Yes" if data.get("isjam") else ""
        data["is_jamchart"] = "Yes" if data.get("isjamchart") else ""
    return data


def fetch_export_rows(
    export_type: str,
    filters: SearchFilters,
    db_path=DB_PATH,
) -> list[dict]:
    conn = get_connection(db_path)
    try:
        if export_type == "performances":
            where, params = build_performance_where(filters)
            sql = PERFORMANCE_SELECT.format(where=where)
            params.append(filters.limit)
            rows = conn.execute(sql, params).fetchall()
        elif export_type == "shows":
            where, params = build_show_where(filters)
            sql = SHOW_SELECT.format(where=where)
            params.append(filters.limit)
            rows = conn.execute(sql, params).fetchall()
        elif export_type == "songs":
            where, params = build_song_where(filters)
            sql = SONG_SELECT.format(where=where)
            params.append(filters.limit)
            rows = conn.execute(sql, params).fetchall()
        else:
            raise ValueError(f"Unknown export type: {export_type}")

        return [_normalize_row(dict(row), export_type) for row in rows]
    finally:
        conn.close()


def columns_for_type(export_type: str) -> list[tuple[str, str]]:
    if export_type == "performances":
        return PERFORMANCE_COLUMNS
    if export_type == "shows":
        return SHOW_COLUMNS
    return SONG_COLUMNS


def export_csv(
    export_type: str,
    filters: SearchFilters,
    output_path: Path | str,
    db_path=DB_PATH,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    rows = fetch_export_rows(export_type, filters, db_path)
    columns = columns_for_type(export_type)

    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow([label for _, label in columns])
        for row in rows:
            writer.writerow([row.get(key, "") for key, _ in columns])

    return output_path


def export_xlsx(
    export_type: str,
    filters: SearchFilters,
    output_path: Path | str,
    db_path=DB_PATH,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    rows = fetch_export_rows(export_type, filters, db_path)
    columns = columns_for_type(export_type)

    wb = Workbook()
    ws = wb.active
    ws.title = export_type.capitalize()

    header_font = Font(bold=True)
    for col_idx, (_, label) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font

    for row_idx, row in enumerate(rows, 2):
        for col_idx, (key, _) in enumerate(columns, 1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))

    for col_idx, (key, _) in enumerate(columns, 1):
        max_len = len(key)
        for row in rows[:200]:
            val = str(row.get(key, "") or "")
            max_len = max(max_len, min(len(val), 50))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 2

    ws.freeze_panes = "A2"
    wb.save(output_path)
    return output_path


def default_export_path(export_type: str, fmt: str) -> Path:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    return DATA_DIR / f"phish_{export_type}.{fmt}"


def get_filter_options(db_path=DB_PATH) -> dict:
    conn = get_connection(db_path)
    try:
        tours = [
            r["tour_name"]
            for r in conn.execute(
                """
                SELECT DISTINCT tour_name FROM shows
                WHERE tour_name IS NOT NULL AND tour_name != ''
                ORDER BY tour_name
                """
            ).fetchall()
        ]
        venues = [
            r["venue"]
            for r in conn.execute(
                """
                SELECT DISTINCT venue FROM shows
                WHERE venue IS NOT NULL AND venue != ''
                ORDER BY venue
                """
            ).fetchall()
        ]
        songs = [
            r["song"]
            for r in conn.execute(
                "SELECT song FROM songs ORDER BY song"
            ).fetchall()
        ]
        return {"tours": tours, "venues": venues, "songs": songs}
    finally:
        conn.close()